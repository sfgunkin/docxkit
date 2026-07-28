#!/usr/bin/env python3
"""
integrate_word_edits.py — Compare a user-edited Word document against the
baseline produced by its generation script.  Reports text, formula (OMML),
and formatting differences at the paragraph/run level.

Usage (baseline mode):
    python integrate_word_edits.py --edited <doc> --script <gen_script> [--output-dir DIR]

Usage (track-changes mode):
    python integrate_word_edits.py --edited <doc> --track-changes [--output-dir DIR]

Output:
    * Stdout table of all changes
    * <stem>_word_edits.json   — machine-readable change list
    * <stem>_user_edited.txt   — full text extract of edited doc
    * <stem>_script_baseline.txt — full text extract of baseline doc (baseline mode only)
"""

from __future__ import annotations

import argparse
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
from copy import deepcopy
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

# Force UTF-8 stdout/stderr on Windows to handle Unicode math symbols
if sys.stdout.encoding and sys.stdout.encoding.lower() not in ("utf-8", "utf8"):
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
if sys.stderr.encoding and sys.stderr.encoding.lower() not in ("utf-8", "utf8"):
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", errors="replace")

from docx import Document
from docx.oxml.ns import qn
from lxml import etree


# ── OMML helpers ─────────────────────────────────────────────────────────

MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def _omml_elements(para_element):
    """Return list of <m:oMath> elements inside a paragraph XML."""
    return para_element.findall(f".//{{{MATH_NS}}}oMath")


def _omml_canonical(elem):
    """Canonical XML string of an OMML element (for comparison)."""
    clone = deepcopy(elem)
    # Strip rsid / revision attributes that don't affect content
    for e in clone.iter():
        for attr in list(e.attrib):
            if "rsid" in attr.lower() or "id" in attr.lower():
                del e.attrib[attr]
    return etree.tostring(clone, method="c14n2").decode()


def _omml_to_unicode(elem) -> str:
    """Best-effort Unicode rendering of OMML for human readability."""
    parts = []
    for node in elem.iter():
        tag = etree.QName(node.tag).localname if isinstance(node.tag, str) else ""
        if tag == "t" and node.text:
            parts.append(node.text)
    return "".join(parts).strip() or "(empty formula)"


# ── Paragraph fingerprinting ────────────────────────────────────────────

def _para_text(para) -> str:
    """Extract all text from a paragraph, including OMML, hyperlinks,
    and Track Changes (include <w:ins> text, exclude <w:del> text)."""
    bits: list[str] = []
    _collect_text(para._element, bits)
    return "".join(bits)


def _collect_text(element, bits: list[str]):
    """Recursively collect text from an XML element, respecting Track Changes.

    Handles: <w:r>, <w:hyperlink>, <w:ins>, <m:oMath>, <m:oMathPara>.
    Skips:   <w:del>, <w:delText>, <w:pPr>, <w:rPr>, <m:oMathParaPr>.
    """
    SKIP_TAGS = {"del", "pPr", "rPr", "sectPr", "oMathParaPr"}
    for child in element:
        tag = etree.QName(child.tag).localname if isinstance(child.tag, str) else ""
        if tag in SKIP_TAGS:
            continue  # skip deleted text and property elements
        elif tag == "r":
            t = child.find(qn("w:t"))
            if t is not None and t.text:
                bits.append(t.text)
        elif tag in ("oMath", "oMathPara"):
            for t_node in child.iter(qn("m:t")):
                if t_node.text:
                    bits.append(t_node.text)
        elif tag == "hyperlink":
            for r in child.findall(qn("w:r")):
                t = r.find(qn("w:t"))
                if t is not None and t.text:
                    bits.append(t.text)
        elif tag == "ins":
            # Track Changes insertion — recurse to collect its runs
            _collect_text(child, bits)
        else:
            # Other wrapper elements (e.g. <w:smartTag>) — recurse
            if len(child) > 0:
                _collect_text(child, bits)


def _normalise(text: str) -> str:
    """Normalise whitespace and typographic variants for matching."""
    text = re.sub(r"\s+", " ", text).strip().lower()
    # Normalise typographic variants that Word may swap
    text = text.replace("\u2212", "-")   # minus sign → hyphen-minus
    text = text.replace("\u2013", "-")   # en dash → hyphen-minus
    text = text.replace("\u2014", "-")   # em dash → hyphen-minus
    text = text.replace("\u2217", "*")   # asterisk operator → asterisk
    text = text.replace("\u2032", "'")   # prime → apostrophe
    text = text.replace("\u2018", "'").replace("\u2019", "'")  # smart quotes
    text = text.replace("\u201c", '"').replace("\u201d", '"')
    return text


def _fingerprint(para) -> str:
    return _normalise(_para_text(para))


# ── Formatting extraction ───────────────────────────────────────────────

def _pt(emu) -> str | None:
    """Convert EMU to points string, or None."""
    if emu is None:
        return None
    return f"{round(emu / 12700, 1)}pt"


def _para_fmt(para) -> dict:
    pf = para.paragraph_format
    return {
        "space_before": _pt(pf.space_before),
        "space_after": _pt(pf.space_after),
        "line_spacing": str(pf.line_spacing) if pf.line_spacing else None,
        "alignment": str(pf.alignment) if pf.alignment is not None else None,
        "first_line_indent": _pt(pf.first_line_indent),
        "left_indent": _pt(pf.left_indent),
    }


def _run_fmt(run) -> dict:
    f = run.font
    return {
        "bold": run.bold,
        "italic": run.italic,
        "underline": run.underline,
        "font_name": f.name,
        "font_size": _pt(f.size) if f.size else None,
        "font_color": str(f.color.rgb) if f.color and f.color.rgb else None,
    }


# ── Core comparison ─────────────────────────────────────────────────────

def _try_close_word_file(filepath: str):
    """Try to close the file in Word via win32com (Windows only)."""
    try:
        import win32com.client
        word = win32com.client.GetObject(Class="Word.Application")
        abspath = os.path.abspath(filepath)
        for doc in word.Documents:
            if os.path.normcase(os.path.abspath(doc.FullName)) == os.path.normcase(abspath):
                doc.Close(SaveChanges=0)
                print(f"  [info] Closed '{filepath}' in Word.", file=sys.stderr)
                return True
    except Exception:
        pass
    return False


def _load_doc(path: str) -> Document:
    """Load a docx, with fallbacks for Word-locked files."""
    path = str(Path(path))  # normalise to OS-native path separators
    try:
        return Document(path)
    except PermissionError:
        print(f"  [warn] PermissionError on '{path}', attempting to close in Word...",
              file=sys.stderr)
        _try_close_word_file(path)
        time.sleep(1)
        try:
            return Document(path)
        except PermissionError:
            print(f"  [warn] Still locked after close attempt: {path}", file=sys.stderr)
    except Exception as exc:
        print(f"  [warn] Could not open '{path}': {exc}", file=sys.stderr)
    # Fallback: copy to temp file (works even if Word holds a lock)
    tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    tmp.close()
    try:
        shutil.copy2(path, tmp.name)
        print(f"  [info] Loaded via temp copy (file locked): {path}", file=sys.stderr)
        return Document(tmp.name)
    finally:
        try:
            os.unlink(tmp.name)
        except OSError:
            pass


def _short_text(text: str, maxlen: int = 12) -> str:
    t = text.strip()
    if len(t) > maxlen:
        return t[:maxlen] + "..."
    return t


def _word_diff(old: str, new: str) -> tuple[str, str]:
    """Return (old_fragment, new_fragment) highlighting word-level diffs."""
    old_words = old.split()
    new_words = new.split()
    sm = SequenceMatcher(None, old_words, new_words)
    old_parts, new_parts = [], []
    for op, i1, i2, j1, j2 in sm.get_opcodes():
        if op == "equal":
            continue
        elif op == "replace":
            old_parts.append(" ".join(old_words[i1:i2]))
            new_parts.append(" ".join(new_words[j1:j2]))
        elif op == "delete":
            old_parts.append(" ".join(old_words[i1:i2]))
        elif op == "insert":
            new_parts.append(" ".join(new_words[j1:j2]))
    old_summary = " | ".join(old_parts) if old_parts else "(none)"
    new_summary = " | ".join(new_parts) if new_parts else "(none)"
    return old_summary, new_summary


def compare_docs(edited_path: str, baseline_path: str) -> list[dict[str, Any]]:
    """Compare edited doc against baseline, return list of change records."""
    doc_edit = _load_doc(edited_path)
    doc_base = _load_doc(baseline_path)

    paras_edit = list(doc_edit.paragraphs)
    paras_base = list(doc_base.paragraphs)

    fp_edit = [_fingerprint(p) for p in paras_edit]
    fp_base = [_fingerprint(p) for p in paras_base]

    # Sequence match on fingerprints
    sm = SequenceMatcher(None, fp_base, fp_edit, autojunk=False)
    changes: list[dict[str, Any]] = []

    for op, i1, i2, j1, j2 in sm.get_opcodes():
        if op == "equal":
            # Check within matched paragraphs for text/formula/format diffs
            for offset in range(i2 - i1):
                bi = i1 + offset
                ei = j1 + offset
                p_base = paras_base[bi]
                p_edit = paras_edit[ei]
                _compare_matched(changes, bi, ei, p_base, p_edit)

        elif op == "replace":
            # Some paragraphs replaced — try to match within the block
            base_block = list(range(i1, i2))
            edit_block = list(range(j1, j2))
            # Attempt pairwise for same-length, otherwise report insert/delete
            paired = min(len(base_block), len(edit_block))
            for k in range(paired):
                bi = base_block[k]
                ei = edit_block[k]
                _compare_matched(changes, bi, ei, paras_base[bi], paras_edit[ei])
            # Extra in base → deleted
            for k in range(paired, len(base_block)):
                bi = base_block[k]
                txt = _para_text(paras_base[bi])
                changes.append({
                    "type": "deleted",
                    "para_base": bi + 1,
                    "para_edit": None,
                    "location": f"¶{bi+1} \"{_short_text(txt)}\"" if txt.strip() else f"¶{bi+1}",
                    "detail": "removed empty paragraph" if not txt.strip()
                             else f"removed: \"{_short_text(txt, 40)}\"",
                    "old": txt,
                    "new": None,
                })
            # Extra in edit → inserted
            for k in range(paired, len(edit_block)):
                ei = edit_block[k]
                txt = _para_text(paras_edit[ei])
                after_base = base_block[-1] if base_block else i1
                changes.append({
                    "type": "inserted",
                    "para_base": None,
                    "para_edit": ei + 1,
                    "location": f"after ¶{after_base+1}",
                    "detail": "blank paragraph (spacing)" if not txt.strip()
                             else f"inserted: \"{_short_text(txt, 40)}\"",
                    "old": None,
                    "new": txt,
                })

        elif op == "delete":
            for bi in range(i1, i2):
                txt = _para_text(paras_base[bi])
                changes.append({
                    "type": "deleted",
                    "para_base": bi + 1,
                    "para_edit": None,
                    "location": f"¶{bi+1}" + (f" \"{_short_text(txt)}\"" if txt.strip() else ""),
                    "detail": "removed empty paragraph" if not txt.strip()
                             else f"removed: \"{_short_text(txt, 40)}\"",
                    "old": txt,
                    "new": None,
                })

        elif op == "insert":
            for ei in range(j1, j2):
                txt = _para_text(paras_edit[ei])
                after_base = i1  # insertion point in baseline
                changes.append({
                    "type": "inserted",
                    "para_base": None,
                    "para_edit": ei + 1,
                    "location": f"after ¶{after_base}",
                    "detail": "blank paragraph (spacing)" if not txt.strip()
                             else f"inserted: \"{_short_text(txt, 40)}\"",
                    "old": None,
                    "new": txt,
                })

    # ── Footnote comparison ────────────────────────────────────────
    fn_changes = _compare_footnotes(edited_path, baseline_path)
    changes.extend(fn_changes)

    return changes


def _extract_footnotes(doc_path: str) -> list[tuple[str, str]]:
    """Extract (fn_id, text) pairs from footnotes.xml, skipping separators."""
    doc = _load_doc(doc_path)
    fn_part = None
    for rel in doc.part.rels.values():
        if 'footnotes' in rel.reltype:
            fn_part = rel.target_part
            break
    if not fn_part:
        return []
    nsmap = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    root = etree.fromstring(fn_part.blob)
    footnotes = []
    for fn in root.findall('.//w:footnote', nsmap):
        fn_id = fn.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}id')
        if fn_id in ('0', '-1'):
            continue
        text = ''
        for t in fn.findall('.//w:t', nsmap):
            text += (t.text or '')
        footnotes.append((fn_id, text.strip()))
    return footnotes


def _compare_footnotes(edited_path: str, baseline_path: str) -> list[dict[str, Any]]:
    """Compare footnotes between edited and baseline documents."""
    fn_edit = _extract_footnotes(edited_path)
    fn_base = _extract_footnotes(baseline_path)

    # Build content-keyed lookups (first 80 chars as key for matching)
    def _key(text: str) -> str:
        return text[:80].strip().lower()

    changes: list[dict[str, Any]] = []

    # Non-empty footnotes only
    edit_nonempty = [(fid, txt) for fid, txt in fn_edit if txt]
    base_nonempty = [(fid, txt) for fid, txt in fn_base if txt]

    # Match by content prefix (first 80 chars) to handle ID renumbering
    edit_keys = {_key(txt): (fid, txt) for fid, txt in edit_nonempty}
    base_keys = {_key(txt): (fid, txt) for fid, txt in base_nonempty}

    # Deleted: in baseline but not in edited
    for k, (fid, txt) in base_keys.items():
        if k not in edit_keys:
            changes.append({
                "type": "footnote_deleted",
                "para_base": f"fn {fid}",
                "para_edit": None,
                "location": f"fn {fid}",
                "detail": f"footnote deleted: \"{_short_text(txt, 60)}\"",
                "old": txt,
                "new": None,
                "substantive": True,
            })

    # Added: in edited but not in baseline
    for k, (fid, txt) in edit_keys.items():
        if k not in base_keys:
            changes.append({
                "type": "footnote_added",
                "para_base": None,
                "para_edit": f"fn {fid}",
                "location": f"fn {fid}",
                "detail": f"footnote added: \"{_short_text(txt, 60)}\"",
                "old": None,
                "new": txt,
                "substantive": True,
            })

    # Changed: same prefix key but different full text
    for k in base_keys:
        if k in edit_keys:
            b_fid, b_txt = base_keys[k]
            e_fid, e_txt = edit_keys[k]
            if b_txt != e_txt:
                changes.append({
                    "type": "footnote_changed",
                    "para_base": f"fn {b_fid}",
                    "para_edit": f"fn {e_fid}",
                    "location": f"fn {b_fid}",
                    "detail": f"footnote text changed",
                    "old": b_txt,
                    "new": e_txt,
                    "substantive": True,
                })

    # Emptied: footnotes that exist in edited but are empty (user deleted content)
    edit_empty = [(fid, txt) for fid, txt in fn_edit if not txt]
    if edit_empty:
        for fid, _ in edit_empty:
            changes.append({
                "type": "footnote_emptied",
                "para_base": None,
                "para_edit": f"fn {fid}",
                "location": f"fn {fid}",
                "detail": f"footnote {fid} emptied (content deleted by user)",
                "old": None,
                "new": "",
                "substantive": True,
            })

    return changes


def _is_inherit_noise(val_base, val_edit) -> bool:
    """True if the diff is just the script being explicit while Word inherits.

    Only suppress when the baseline (script) sets an explicit value and the
    edited (user) doc shows None (inherited from style).  The reverse case
    — script has None and user sets an explicit value — is a real edit.
    """
    return val_base is not None and val_edit is None


def _para_has_hyperlinks(para) -> bool:
    """True if the paragraph XML contains any w:hyperlink elements."""
    return len(para._element.findall(qn("w:hyperlink"))) > 0


def _is_hyperlink_noise(para, prop: str, old_val, new_val) -> bool:
    """True if a run_format change is hyperlink styling bleed.

    When make_hyperlink() creates blue underlined text, Word often
    propagates underline/font_color/italic to adjacent runs in the
    saved XML. These are not real user edits.
    """
    if not _para_has_hyperlinks(para):
        return False
    # underline: None → True on a paragraph with hyperlinks
    if prop == "underline" and old_val is None and new_val is True:
        return True
    # font_color: None → any hex value
    if prop == "font_color" and old_val is None and new_val is not None:
        return True
    # italic: None → True can also bleed from adjacent hyperlink styling
    if prop == "italic" and old_val is None and new_val is True:
        return True
    return False


def _compare_matched(changes, bi, ei, p_base, p_edit):
    """Compare two matched paragraphs for text, formula, and format diffs."""
    text_base = _para_text(p_base)
    text_edit = _para_text(p_edit)
    loc_label = f"¶{bi+1} \"{_short_text(text_base)}\"" if text_base.strip() else f"¶{bi+1}"

    # --- Text diff ---
    if _normalise(text_base) != _normalise(text_edit):
        # More specific: word-level diff summary
        old_frag, new_frag = _word_diff(text_base, text_edit)
        changes.append({
            "type": "text",
            "para_base": bi + 1,
            "para_edit": ei + 1,
            "location": loc_label,
            "detail": f"\"{_short_text(old_frag, 30)}\" → \"{_short_text(new_frag, 30)}\"",
            "old": text_base,
            "new": text_edit,
        })

    # --- Formula (OMML) diff ---
    omml_base = _omml_elements(p_base._element)
    omml_edit = _omml_elements(p_edit._element)

    # Compare by index
    n = max(len(omml_base), len(omml_edit))
    for k in range(n):
        ob = omml_base[k] if k < len(omml_base) else None
        oe = omml_edit[k] if k < len(omml_edit) else None
        if ob is not None and oe is not None:
            if _omml_canonical(ob) != _omml_canonical(oe):
                # Skip if Unicode rendering is identical (XML-only difference)
                u_old = _omml_to_unicode(ob)
                u_new = _omml_to_unicode(oe)
                if _normalise(u_old) == _normalise(u_new):
                    continue  # same formula, different XML or typographic variants — noise
                changes.append({
                    "type": "formula",
                    "para_base": bi + 1,
                    "para_edit": ei + 1,
                    "location": loc_label,
                    "detail": f"{u_old} → {u_new}",
                    "old_omml": _omml_canonical(ob),
                    "new_omml": _omml_canonical(oe),
                    "old_unicode": u_old,
                    "new_unicode": u_new,
                })
        elif ob is None and oe is not None:
            changes.append({
                "type": "formula",
                "para_base": bi + 1,
                "para_edit": ei + 1,
                "location": loc_label,
                "detail": f"added formula: {_omml_to_unicode(oe)}",
                "old_omml": None,
                "new_omml": _omml_canonical(oe),
                "old_unicode": None,
                "new_unicode": _omml_to_unicode(oe),
            })
        elif ob is not None and oe is None:
            changes.append({
                "type": "formula",
                "para_base": bi + 1,
                "para_edit": ei + 1,
                "location": loc_label,
                "detail": f"removed formula: {_omml_to_unicode(ob)}",
                "old_omml": _omml_canonical(ob),
                "new_omml": None,
                "old_unicode": _omml_to_unicode(ob),
                "new_unicode": None,
            })

    # --- Paragraph format diff ---
    fmt_base = _para_fmt(p_base)
    fmt_edit = _para_fmt(p_edit)
    for key in fmt_base:
        vb = fmt_base[key]
        ve = fmt_edit[key]
        if vb != ve and not _is_inherit_noise(vb, ve):
            changes.append({
                "type": "format",
                "para_base": bi + 1,
                "para_edit": ei + 1,
                "location": loc_label,
                "detail": f"{key}: {vb} → {ve}",
                "property": key,
                "old_value": vb,
                "new_value": ve,
            })

    # --- Run-level format diff ---
    runs_base = list(p_base.runs)
    runs_edit = list(p_edit.runs)
    # Match runs by index (best effort)
    n_runs = min(len(runs_base), len(runs_edit))
    for k in range(n_runs):
        rb = runs_base[k]
        re_ = runs_edit[k]
        fb = _run_fmt(rb)
        fe = _run_fmt(re_)
        for key in fb:
            vb = fb[key]
            ve = fe[key]
            if vb != ve and not _is_inherit_noise(vb, ve) \
                    and not _is_hyperlink_noise(p_edit, key, vb, ve):
                run_text = (rb.text or "") or (re_.text or "")
                changes.append({
                    "type": "run_format",
                    "para_base": bi + 1,
                    "para_edit": ei + 1,
                    "location": loc_label,
                    "detail": f"\"{_short_text(run_text)}\" {key}: {vb} → {ve}",
                    "run_index": k,
                    "run_text": run_text,
                    "property": key,
                    "old_value": str(vb),
                    "new_value": str(ve),
                })


# ── Baseline generation ─────────────────────────────────────────────────

def _generate_baseline(script_path: str, output_dir: str, edited_path: str) -> str:
    """Run the generation script and return path to the generated docx.

    Always snapshots the edited file before running the script in case
    the script saves to the same path.
    """
    script = Path(script_path)
    edited = Path(edited_path).resolve()

    # Always snapshot — the script may overwrite the edited file
    snapshot = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    snapshot.close()
    shutil.copy2(str(edited), snapshot.name)
    edited_mtime = edited.stat().st_mtime

    # Run the script
    print(f"  [info] Running generation script: {script_path}", file=sys.stderr)
    result = subprocess.run(
        [sys.executable, str(script)],
        cwd=str(script.parent),
        capture_output=True, encoding="utf-8", errors="replace", timeout=300,
    )
    if result.returncode != 0:
        shutil.copy2(snapshot.name, str(edited))
        os.unlink(snapshot.name)
        print(f"  [error] Script failed:\n{result.stderr[:2000]}", file=sys.stderr)
        sys.exit(1)

    # Find the generated baseline: most recently modified .docx
    search_dirs = [script.parent, script.parent.parent / "Documents"]
    recent = None
    recent_mtime = 0
    for d in search_dirs:
        if d.is_dir():
            for f in d.glob("*.docx"):
                mt = f.stat().st_mtime
                if mt > recent_mtime:
                    recent = f
                    recent_mtime = mt

    if not recent:
        shutil.copy2(snapshot.name, str(edited))
        os.unlink(snapshot.name)
        print("  [error] Could not find generated docx after running script.",
              file=sys.stderr)
        sys.exit(1)

    baseline_path = str(recent.resolve())

    # If the script overwrote the edited file, move the generated baseline
    # to a temp location and restore the user's version.
    if recent.resolve() == edited or edited.stat().st_mtime > edited_mtime:
        baseline_tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
        baseline_tmp.close()
        shutil.copy2(baseline_path, baseline_tmp.name)
        shutil.copy2(snapshot.name, str(edited))
        baseline_path = baseline_tmp.name
        print("  [info] Script overwrote edited file; restored from snapshot.",
              file=sys.stderr)

    os.unlink(snapshot.name)
    return baseline_path


# ── Text extraction ─────────────────────────────────────────────────────

def _extract_text(doc_path: str) -> str:
    doc = _load_doc(doc_path)
    lines = []
    for i, p in enumerate(doc.paragraphs, 1):
        txt = _para_text(p)
        lines.append(f"[¶{i}] {txt}")
    return "\n".join(lines)


# ── Output formatting ───────────────────────────────────────────────────

def _has_track_changes(doc_path: str) -> bool:
    """Check if the document contains unresolved Track Changes."""
    doc = _load_doc(doc_path)
    for p in doc.paragraphs:
        for child in p._element.iter():
            tag = etree.QName(child.tag).localname if isinstance(child.tag, str) else ""
            if tag in ("ins", "del"):
                return True
    return False


# ── Track Changes extraction ─────────────────────────────────────────────

WML_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _tc_attr(elem, attr: str) -> str | None:
    """Read a w:-namespaced attribute from a Track Changes element."""
    return elem.get(qn(f"w:{attr}")) or elem.get(attr)


def _tc_text_from_runs(parent) -> str:
    """Collect text from <w:r>/<w:t> children (used for <w:ins> content)."""
    parts: list[str] = []
    for r in parent.iter(qn("w:r")):
        t = r.find(qn("w:t"))
        if t is not None and t.text:
            parts.append(t.text)
    return "".join(parts)


def _tc_deltext(parent) -> str:
    """Collect text from <w:delText> children (used for <w:del> content)."""
    parts: list[str] = []
    for dt in parent.iter(qn("w:delText")):
        if dt.text:
            parts.append(dt.text)
    return "".join(parts)


def _tc_snippet(full_text: str, changed_text: str, maxlen: int = 60) -> str:
    """Build a context snippet: the full paragraph text with the changed portion highlighted."""
    if not full_text.strip():
        return "(empty paragraph)"
    if not changed_text.strip():
        return _short_text(full_text, maxlen)
    # Try to find the changed text in the full text for context
    idx = full_text.find(changed_text)
    if idx >= 0:
        before = full_text[max(0, idx - 20):idx]
        after = full_text[idx + len(changed_text):idx + len(changed_text) + 20]
        snippet = f"...{before}[{changed_text}]{after}..."
        return _short_text(snippet, maxlen)
    return _short_text(full_text, maxlen)


def _describe_rpr_change(rpr_change) -> list[str]:
    """Describe what changed in a <w:rPrChange> element by comparing old vs current rPr."""
    diffs = []
    # The rPrChange contains the OLD formatting; the parent <w:rPr> has the NEW
    parent_rpr = rpr_change.getparent()
    old_rpr = rpr_change

    for prop_tag, label in [("b", "bold"), ("i", "italic"), ("u", "underline"),
                            ("strike", "strikethrough")]:
        old_has = old_rpr.find(qn(f"w:{prop_tag}")) is not None
        new_has = parent_rpr.find(qn(f"w:{prop_tag}")) is not None
        if old_has != new_has:
            if new_has:
                diffs.append(f"+{label}")
            else:
                diffs.append(f"-{label}")

    # Font size
    old_sz = old_rpr.find(qn("w:sz"))
    new_sz = parent_rpr.find(qn("w:sz"))
    old_val = old_sz.get(qn("w:val")) if old_sz is not None else None
    new_val = new_sz.get(qn("w:val")) if new_sz is not None else None
    if old_val != new_val:
        # Sizes are in half-points
        old_pt = f"{int(old_val)/2}pt" if old_val else "inherited"
        new_pt = f"{int(new_val)/2}pt" if new_val else "inherited"
        diffs.append(f"size: {old_pt} → {new_pt}")

    # Font name
    old_fonts = old_rpr.find(qn("w:rFonts"))
    new_fonts = parent_rpr.find(qn("w:rFonts"))
    old_font = old_fonts.get(qn("w:ascii")) if old_fonts is not None else None
    new_font = new_fonts.get(qn("w:ascii")) if new_fonts is not None else None
    if old_font != new_font:
        diffs.append(f"font: {old_font or 'inherited'} → {new_font or 'inherited'}")

    return diffs if diffs else ["formatting changed"]


def _describe_ppr_change(ppr_change) -> list[str]:
    """Describe what changed in a <w:pPrChange> element."""
    diffs = []
    parent_ppr = ppr_change.getparent()
    old_ppr = ppr_change

    # Alignment / justification
    old_jc = old_ppr.find(qn("w:jc"))
    new_jc = parent_ppr.find(qn("w:jc"))
    old_val = old_jc.get(qn("w:val")) if old_jc is not None else None
    new_val = new_jc.get(qn("w:val")) if new_jc is not None else None
    if old_val != new_val:
        diffs.append(f"alignment: {old_val or 'default'} → {new_val or 'default'}")

    # Spacing
    old_sp = old_ppr.find(qn("w:spacing"))
    new_sp = parent_ppr.find(qn("w:spacing"))
    for attr_name in ["before", "after", "line"]:
        old_v = old_sp.get(qn(f"w:{attr_name}")) if old_sp is not None else None
        new_v = new_sp.get(qn(f"w:{attr_name}")) if new_sp is not None else None
        if old_v != new_v:
            diffs.append(f"spacing_{attr_name}: {old_v or 'default'} → {new_v or 'default'}")

    # Indentation
    old_ind = old_ppr.find(qn("w:ind"))
    new_ind = parent_ppr.find(qn("w:ind"))
    for attr_name in ["left", "right", "firstLine", "hanging"]:
        old_v = old_ind.get(qn(f"w:{attr_name}")) if old_ind is not None else None
        new_v = new_ind.get(qn(f"w:{attr_name}")) if new_ind is not None else None
        if old_v != new_v:
            diffs.append(f"indent_{attr_name}: {old_v or 'default'} → {new_v or 'default'}")

    return diffs if diffs else ["paragraph formatting changed"]


def extract_track_changes(doc_path: str) -> list[dict[str, Any]]:
    """Extract all Track Changes from a Word document.

    Returns a list of change records with types:
      tc_insert, tc_delete, tc_move, tc_format
    Each record includes author, date, paragraph index, context, and detail.
    """
    doc = _load_doc(doc_path)
    changes: list[dict[str, Any]] = []

    # First pass: collect moveFrom/moveTo IDs so we can pair them
    move_from_ids: dict[str, int] = {}   # id -> para index (1-based)
    move_to_ids: dict[str, int] = {}
    move_from_text: dict[str, str] = {}

    for pi, para in enumerate(doc.paragraphs):
        elem = para._element
        for mf in elem.iter(qn("w:moveFrom")):
            mid = _tc_attr(mf, "id") or ""
            if mid:
                move_from_ids[mid] = pi + 1
                move_from_text[mid] = _tc_deltext(mf) or _tc_text_from_runs(mf)
        for mt in elem.iter(qn("w:moveTo")):
            mid = _tc_attr(mt, "id") or ""
            if mid:
                move_to_ids[mid] = pi + 1

    # Build set of IDs that have both from+to (true moves)
    paired_move_ids = set(move_from_ids.keys()) & set(move_to_ids.keys())
    # Track which moveFrom/moveTo we've already reported
    reported_moves: set[str] = set()

    # Second pass: walk paragraphs and extract changes
    for pi, para in enumerate(doc.paragraphs):
        elem = para._element
        para_idx = pi + 1
        full_text = _para_text(para)
        loc_label = f"¶{para_idx}" + (f" \"{_short_text(full_text)}\"" if full_text.strip() else "")

        for child in elem.iter():
            tag = etree.QName(child.tag).localname if isinstance(child.tag, str) else ""

            if tag == "ins":
                # Skip if this is inside a moveTo (handled separately)
                parent_tags = set()
                p = child.getparent()
                while p is not None:
                    pt = etree.QName(p.tag).localname if isinstance(p.tag, str) else ""
                    parent_tags.add(pt)
                    p = p.getparent()
                if "moveTo" in parent_tags:
                    continue

                text = _tc_text_from_runs(child)
                if not text.strip():
                    continue
                author = _tc_attr(child, "author") or "unknown"
                date = _tc_attr(child, "date") or ""
                changes.append({
                    "type": "tc_insert",
                    "para_edit": para_idx,
                    "location": loc_label,
                    "detail": f"inserted: \"{_short_text(text, 40)}\"",
                    "text": text,
                    "context": _tc_snippet(full_text, text),
                    "author": author,
                    "date": date,
                })

            elif tag == "del":
                # Skip if this is inside a moveFrom (handled separately)
                parent_tags = set()
                p = child.getparent()
                while p is not None:
                    pt = etree.QName(p.tag).localname if isinstance(p.tag, str) else ""
                    parent_tags.add(pt)
                    p = p.getparent()
                if "moveFrom" in parent_tags:
                    continue

                text = _tc_deltext(child)
                if not text.strip():
                    continue
                author = _tc_attr(child, "author") or "unknown"
                date = _tc_attr(child, "date") or ""
                changes.append({
                    "type": "tc_delete",
                    "para_edit": para_idx,
                    "location": loc_label,
                    "detail": f"deleted: \"{_short_text(text, 40)}\"",
                    "text": text,
                    "context": _tc_snippet(full_text, text),
                    "author": author,
                    "date": date,
                })

            elif tag == "moveTo":
                mid = _tc_attr(child, "id") or ""
                if mid in paired_move_ids and mid not in reported_moves:
                    reported_moves.add(mid)
                    text = move_from_text.get(mid, "")
                    author = _tc_attr(child, "author") or "unknown"
                    date = _tc_attr(child, "date") or ""
                    from_para = move_from_ids.get(mid, 0)
                    to_para = move_to_ids.get(mid, 0)
                    changes.append({
                        "type": "tc_move",
                        "para_edit": para_idx,
                        "location": f"¶{from_para} → ¶{to_para}",
                        "detail": f"moved: \"{_short_text(text, 40)}\"",
                        "text": text,
                        "from_para": from_para,
                        "to_para": to_para,
                        "author": author,
                        "date": date,
                    })

            elif tag == "moveFrom":
                mid = _tc_attr(child, "id") or ""
                if mid not in paired_move_ids:
                    # Orphaned moveFrom — treat as delete
                    text = _tc_deltext(child) or _tc_text_from_runs(child)
                    if not text.strip():
                        continue
                    author = _tc_attr(child, "author") or "unknown"
                    date = _tc_attr(child, "date") or ""
                    changes.append({
                        "type": "tc_delete",
                        "para_edit": para_idx,
                        "location": loc_label,
                        "detail": f"deleted (orphan move): \"{_short_text(text, 40)}\"",
                        "text": text,
                        "context": _tc_snippet(full_text, text),
                        "author": author,
                        "date": date,
                    })

            elif tag == "rPrChange":
                author = _tc_attr(child, "author") or "unknown"
                date = _tc_attr(child, "date") or ""
                # Get the run text for context
                run_elem = child.getparent()  # <w:rPr>
                if run_elem is not None:
                    run_elem = run_elem.getparent()  # <w:r>
                run_text = ""
                if run_elem is not None:
                    t = run_elem.find(qn("w:t"))
                    if t is not None and t.text:
                        run_text = t.text

                fmt_diffs = _describe_rpr_change(child)
                changes.append({
                    "type": "tc_format",
                    "para_edit": para_idx,
                    "location": loc_label,
                    "detail": f"\"{_short_text(run_text)}\" {', '.join(fmt_diffs)}",
                    "run_text": run_text,
                    "format_changes": fmt_diffs,
                    "author": author,
                    "date": date,
                })

            elif tag == "pPrChange":
                author = _tc_attr(child, "author") or "unknown"
                date = _tc_attr(child, "date") or ""
                fmt_diffs = _describe_ppr_change(child)
                changes.append({
                    "type": "tc_format",
                    "para_edit": para_idx,
                    "location": loc_label,
                    "detail": f"para format: {', '.join(fmt_diffs)}",
                    "format_changes": fmt_diffs,
                    "author": author,
                    "date": date,
                })

    return changes


# ── Output formatting ───────────────────────────────────────────────────


def _print_tc_table(changes: list[dict]):
    """Print a table formatted for track-changes output (includes author/date columns)."""
    if not changes:
        print("\nNo tracked changes found in the document.")
        return

    print(f"\nFound {len(changes)} tracked change(s):\n")
    print(f"{'#':>3}  | {'Type':<11} | {'Author':<16} | {'Location':<18} | Detail")
    print(f"{'---':>3}--|{'-'*12}-|{'-'*17}-|{'-'*19}-|{'-'*50}")

    for i, c in enumerate(changes, 1):
        typ = c["type"]
        author = c.get("author", "")
        loc = c.get("location", "")
        detail = c.get("detail", "")
        if len(author) > 16:
            author = author[:13] + "..."
        if len(loc) > 18:
            loc = loc[:15] + "..."
        if len(detail) > 50:
            detail = detail[:47] + "..."
        print(f"{i:>3}  | {typ:<11} | {author:<16} | {loc:<18} | {detail}")


def _print_table(changes: list[dict], track_changes: bool = False,
                  show_all: bool = False):
    if not changes:
        print("\nNo differences found between user edits and script baseline.")
        return

    if track_changes:
        print("\n  [note] Document contains Track Changes (ins/del). "
              "Inserted text is included; deleted text is excluded.")

    # Split substantive vs formatting-only
    SUBSTANTIVE_TYPES = {"text", "formula", "format", "deleted", "inserted",
                         "footnote_deleted", "footnote_added", "footnote_changed",
                         "footnote_emptied"}
    substantive = [c for c in changes if c["type"] in SUBSTANTIVE_TYPES]
    formatting = [c for c in changes if c["type"] not in SUBSTANTIVE_TYPES]

    to_show = changes if show_all else substantive
    hidden = len(formatting) if not show_all else 0

    if not to_show and hidden:
        print(f"\nNo substantive changes. ({hidden} run_format changes hidden; use --all to see)")
        return

    n_sub = len(substantive)
    n_fmt = len(formatting)
    label = f"{n_sub} substantive"
    if n_fmt:
        label += f" + {n_fmt} formatting"
        if not show_all:
            label += " (hidden)"
    print(f"\nFound {label} change(s):\n")

    # Header
    print(f"{'#':>3}  | {'Type':<11} | {'Para':<5} | Detail")
    print(f"{'---':>3}--|{'-'*12}-|{'-'*6}-|{'-'*60}")

    for i, c in enumerate(to_show, 1):
        typ = c["type"]
        para = c.get("para_base") or c.get("para_edit") or ""
        detail = c.get("detail", "")
        if len(detail) > 60:
            detail = detail[:57] + "..."
        # Footnote changes use "fn X" not "¶X"
        para_str = str(para)
        if para_str.startswith("fn "):
            print(f"{i:>3}  | {typ:<11} | {para_str:<5} | {detail}")
        else:
            print(f"{i:>3}  | {typ:<11} | ¶{para:<4} | {detail}")

        # For text changes, print old/new excerpts inline
        if typ == "text" and "old" in c and "new" in c:
            old_t = c["old"]
            new_t = c["new"]
            # Skip timestamp-only diffs
            if old_t[:10] == new_t[:10] and "—" in old_t[:30]:
                continue
            # Show word-level diff excerpt (up to 200 chars each)
            old_frag, new_frag = _word_diff(old_t, new_t)
            if len(old_frag) > 200:
                old_frag = old_frag[:197] + "..."
            if len(new_frag) > 200:
                new_frag = new_frag[:197] + "..."
            print(f"       OLD: {old_frag}")
            print(f"       NEW: {new_frag}")
        elif typ == "deleted" and "old" in c:
            excerpt = c["old"][:120]
            if len(c["old"]) > 120:
                excerpt += "..."
            print(f"       DEL: {excerpt}")
        elif typ == "footnote_deleted" and "old" in c:
            excerpt = c["old"][:120]
            if len(c["old"]) > 120:
                excerpt += "..."
            print(f"       DEL: {excerpt}")
        elif typ == "footnote_added" and "new" in c:
            excerpt = c["new"][:120]
            if len(c["new"]) > 120:
                excerpt += "..."
            print(f"       ADD: {excerpt}")
        elif typ == "footnote_changed" and "old" in c and "new" in c:
            old_frag, new_frag = _word_diff(c["old"], c["new"])
            if len(old_frag) > 200:
                old_frag = old_frag[:197] + "..."
            if len(new_frag) > 200:
                new_frag = new_frag[:197] + "..."
            print(f"       OLD: {old_frag}")
            print(f"       NEW: {new_frag}")
        elif typ == "footnote_emptied":
            print(f"       (user deleted footnote content)")


def _save_json(changes: list[dict], path: str):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(changes, f, indent=2, ensure_ascii=False)


# ── Word-native comparison (fast mode) ────────────────────────────────────

def _compare_via_word(baseline_path: str, edited_path: str,
                      out_dir: str) -> tuple[str, list[dict[str, Any]]]:
    """Use Word's CompareDocuments for fast native comparison.

    Returns (redline_path, changes) where changes is the same format
    as extract_track_changes().
    """
    import win32com.client

    baseline_abs = os.path.abspath(baseline_path)
    edited_abs = os.path.abspath(edited_path)

    redline_path = os.path.join(
        out_dir,
        Path(edited_path).stem + "_redline.docx",
    )

    word = None
    try:
        # Try to connect to existing Word instance first
        try:
            word = win32com.client.GetObject(Class="Word.Application")
            reused = True
        except Exception:
            word = win32com.client.gencache.EnsureDispatch("Word.Application")
            reused = False

        word.Visible = False
        word.DisplayAlerts = 0  # wdAlertsNone

        # Open both docs (not ReadOnly — we need to accept revisions)
        doc_base = word.Documents.Open(baseline_abs)
        doc_edit = word.Documents.Open(edited_abs)

        # Accept all existing track changes so CompareDocuments doesn't
        # prompt about pre-existing revisions
        if doc_base.Revisions.Count > 0:
            doc_base.AcceptAllRevisions()
        if doc_edit.Revisions.Count > 0:
            doc_edit.AcceptAllRevisions()

        # CompareDocuments parameters:
        #   Destination=2 (wdCompareDestinationNew)
        #   Granularity=1 (wdGranularityWordLevel)
        #   CompareFormatting=True
        #   CompareCaseChanges=True
        #   CompareWhitespace=False  (ignore whitespace-only diffs)
        redline = word.CompareDocuments(
            doc_base, doc_edit,
            Destination=2,
            Granularity=1,
            CompareFormatting=True,
            CompareCaseChanges=True,
            CompareWhitespace=False,
            CompareTables=True,
            CompareHeaders=True,
            CompareFootnotes=True,
            CompareFields=False,  # skip field code differences
        )

        # Save redline document
        # wdFormatDocumentDefault = 16 (.docx)
        redline.SaveAs2(os.path.abspath(redline_path), FileFormat=16)
        redline.Close(SaveChanges=0)
        doc_edit.Close(SaveChanges=0)
        doc_base.Close(SaveChanges=0)

        if not reused:
            word.Quit()
            word = None

    except Exception as exc:
        # Clean up Word if we started it
        if word is not None:
            try:
                word.Quit()
            except Exception:
                pass
        raise RuntimeError(f"Word CompareDocuments failed: {exc}") from exc

    print(f"  Redline saved: {redline_path}", file=sys.stderr)

    # Parse the redline document for track changes
    changes = extract_track_changes(redline_path)
    return redline_path, changes


# ── Main ─────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(
        description="Compare user-edited Word doc against script-generated baseline.")
    parser.add_argument("--edited", required=True, help="Path to user-edited .docx")
    parser.add_argument("--script", default=None, help="Path to generation script (.py)")
    parser.add_argument("--baseline", default=None,
                        help="Path to pre-generated baseline .docx (skip re-generation)")
    parser.add_argument("--track-changes", action="store_true",
                        help="Extract changes from Track Changes markup (no baseline needed)")
    parser.add_argument("--output-dir", default=None,
                        help="Directory for output files (default: same as edited doc)")
    parser.add_argument("--all", action="store_true",
                        help="Show all changes including run_format noise (default: substantive only)")
    parser.add_argument("--fast", action="store_true",
                        help="Use Word's native CompareDocuments (requires Word + --baseline)")
    args = parser.parse_args()

    if not args.track_changes and not args.script and not args.fast:
        parser.error("--script is required unless --track-changes or --fast is set")
    if args.fast and not args.baseline:
        parser.error("--fast requires --baseline (path to previously generated .docx)")

    edited_path = os.path.abspath(args.edited)

    if not os.path.isfile(edited_path):
        print(f"Error: edited doc not found: {edited_path}", file=sys.stderr)
        sys.exit(1)

    out_dir = args.output_dir or os.path.dirname(edited_path)
    stem = Path(edited_path).stem

    # ── Track Changes mode ──────────────────────────────────────────
    if args.track_changes:
        print(f"  Edited: {edited_path}", file=sys.stderr)
        print("Extracting tracked changes...", file=sys.stderr)
        changes = extract_track_changes(edited_path)

        # Output table
        _print_tc_table(changes)

        # Save JSON
        json_path = os.path.join(out_dir, f"{stem}_word_edits.json")
        _save_json(changes, json_path)
        print(f"\n  JSON saved: {json_path}", file=sys.stderr)

        # Save text extract (accepted-state)
        txt_edited = os.path.join(out_dir, f"{stem}_user_edited.txt")
        with open(txt_edited, "w", encoding="utf-8") as f:
            f.write(_extract_text(edited_path))
        print(f"  Text extract saved: {txt_edited}", file=sys.stderr)
        return

    # ── Fast mode (Word-native CompareDocuments) ─────────────────────
    if args.fast:
        baseline_path = os.path.abspath(args.baseline)
        if not os.path.isfile(baseline_path):
            print(f"Error: baseline doc not found: {baseline_path}", file=sys.stderr)
            sys.exit(1)
        if os.path.normcase(baseline_path) == os.path.normcase(edited_path):
            print("  [error] Baseline and edited are the same file.", file=sys.stderr)
            sys.exit(1)

        print(f"  Baseline: {baseline_path}", file=sys.stderr)
        print(f"  Edited:   {edited_path}", file=sys.stderr)
        print("Comparing via Word CompareDocuments (fast)...", file=sys.stderr)
        t0 = time.time()
        redline_path, changes = _compare_via_word(baseline_path, edited_path, out_dir)
        elapsed = time.time() - t0
        print(f"  Comparison completed in {elapsed:.1f}s", file=sys.stderr)

        # Output table (track-changes format)
        _print_tc_table(changes)

        # Save JSON
        json_path = os.path.join(out_dir, f"{stem}_word_edits.json")
        _save_json(changes, json_path)
        print(f"\n  JSON saved: {json_path}", file=sys.stderr)

        # Save text extracts
        txt_edited = os.path.join(out_dir, f"{stem}_user_edited.txt")
        with open(txt_edited, "w", encoding="utf-8") as f:
            f.write(_extract_text(edited_path))
        print(f"  Text extract saved: {txt_edited}", file=sys.stderr)
        return

    # ── Baseline comparison mode (original) ─────────────────────────
    script_path = os.path.abspath(args.script)
    if not os.path.isfile(script_path):
        print(f"Error: script not found: {script_path}", file=sys.stderr)
        sys.exit(1)

    # Step 1: Generate or use baseline
    if args.baseline:
        baseline_path = os.path.abspath(args.baseline)
        if not os.path.isfile(baseline_path):
            print(f"Error: baseline doc not found: {baseline_path}", file=sys.stderr)
            sys.exit(1)
    else:
        print("Generating baseline document from script...", file=sys.stderr)
        baseline_path = _generate_baseline(script_path, out_dir, edited_path)
    print(f"  Baseline: {baseline_path}", file=sys.stderr)
    print(f"  Edited:   {edited_path}", file=sys.stderr)

    # Safety check: edited and baseline must be different files
    if os.path.normcase(os.path.abspath(baseline_path)) == \
       os.path.normcase(os.path.abspath(edited_path)):
        print("  [error] Baseline and edited paths resolve to the same file.\n"
              "          The generation script likely saves to the same path.\n"
              "          Use --baseline to provide a separate baseline, or check\n"
              "          the script's output path.", file=sys.stderr)
        sys.exit(1)

    # Step 2: Check for Track Changes
    has_tc = _has_track_changes(edited_path)
    if has_tc:
        print("  [info] Track Changes detected in edited document.", file=sys.stderr)

    # Step 3: Compare
    print("Comparing documents...", file=sys.stderr)
    changes = compare_docs(edited_path, baseline_path)

    # Step 4: Output table
    _print_table(changes, track_changes=has_tc, show_all=args.all)

    # Step 5: Save JSON
    json_path = os.path.join(out_dir, f"{stem}_word_edits.json")
    _save_json(changes, json_path)
    print(f"\n  JSON saved: {json_path}", file=sys.stderr)

    # Step 6: Save text extracts
    txt_edited = os.path.join(out_dir, f"{stem}_user_edited.txt")
    txt_baseline = os.path.join(out_dir, f"{stem}_script_baseline.txt")
    with open(txt_edited, "w", encoding="utf-8") as f:
        f.write(_extract_text(edited_path))
    with open(txt_baseline, "w", encoding="utf-8") as f:
        f.write(_extract_text(baseline_path))
    print(f"  Text extracts saved: {txt_edited}, {txt_baseline}", file=sys.stderr)


if __name__ == "__main__":
    main()
